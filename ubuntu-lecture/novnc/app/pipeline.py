import json
import os
import subprocess
import sys
from pathlib import Path

WHISPER_MODEL = os.environ.get('WHISPER_MODEL', 'medium')
COURSE_XLSX = os.environ.get('COURSE_XLSX', '')

_model = None


def lookup_course(no):
    """연번으로 과정 정보를 찾는다. 엑셀이 없거나 연번이 없으면 {} 를 돌려준다."""
    if not no or not COURSE_XLSX or not Path(COURSE_XLSX).is_file():
        return {}
    import openpyxl
    ws = openpyxl.load_workbook(COURSE_XLSX, read_only=True)['콘텐츠리스트']
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    col = {name: i for i, name in enumerate(header)}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if str(row[0]) == str(no):
            pick = lambda k: row[col[k]] if k in col else None
            desc = (pick('주요훈련내용') or '').replace('_x000B_', '\n')
            return {
                'title': pick('콘텐츠명'),
                'minutes': pick('훈련시간(분)'),
                'desc': desc,
                'grade': pick('AI추천'),
                'reason': pick('추천사유'),
            }
    return {}


def _stamp(sec):
    m, s = divmod(int(sec), 60)
    h, m = divmod(m, 60)
    return f'{h:d}:{m:02d}:{s:02d}' if h else f'{m:02d}:{s:02d}'


def transcribe(note):
    global _model
    if _model is None:
        from faster_whisper import WhisperModel
        print('모델로딩', flush=True)
        _model = WhisperModel(WHISPER_MODEL, device='cpu', compute_type='int8')
    print('전사', flush=True)
    segments, _ = _model.transcribe(str(note / 'audio.ogg'), language='ko', vad_filter=True, beam_size=5)
    lines = [f'[{_stamp(s.start)}] {s.text.strip()}' for s in segments]
    (note / 'transcript.txt').write_text('\n'.join(lines) + '\n', encoding='utf-8')
    return len(lines)


PROMPT = """너는 사내 원격훈련 강의의 개인 복습 노트를 만든다.
stdin 은 강의 음성을 자동 전사한 텍스트라 인식 오류가 섞여 있다. 문맥으로 바로잡되 강의에 없는 내용은 지어내지 마라.
도구를 쓰지 말고 마크다운 본문만 출력한다.

과정: {title}
주요훈련내용: {desc}
수강 목적: {reason}

형식:
# {title}
## 한 줄 요약
## 핵심 내용
- 5~10개. 각 불릿 끝에 근거 시각 [mm:ss]
## 용어 정리
| 용어 | 뜻 | 강의 속 예시 |
## 수강 목적과 연결
- 수강 목적이 비어 있으면 이 절을 생략한다. 있으면 강의 내용 중 바로 적용할 점 3개 이내
## 복습 질문
- 5개. 각 질문 아래 <details><summary>답</summary>...</details>
"""


def summarize(note):
    meta = json.loads((note / 'meta.json').read_text(encoding='utf-8'))
    prompt = PROMPT.format(
        title=meta.get('title') or note.name,
        desc=(meta.get('desc') or '').strip() or '없음',
        reason=meta.get('reason') or '',
    )
    print('요약', flush=True)
    out = subprocess.run(
        ['claude', '-p', prompt, '--output-format', 'text'],
        input=(note / 'transcript.txt').read_text(encoding='utf-8'),
        capture_output=True, text=True, timeout=1800, cwd=note,
    )
    if out.returncode != 0 or not out.stdout.strip():
        raise RuntimeError(out.stderr.strip() or out.stdout.strip() or 'claude 출력 없음')
    (note / 'summary.md').write_text(out.stdout, encoding='utf-8')


def process(note):
    """transcript.txt·summary.md 가 이미 있으면 그 단계는 건너뛴다. 실패하면 error.log 를 남긴다."""
    note = Path(note)
    err = note / 'error.log'
    err.unlink(missing_ok=True)
    try:
        if not (note / 'transcript.txt').exists():
            transcribe(note)
        if not (note / 'summary.md').exists():
            summarize(note)
        print('완료', flush=True)
    except Exception as e:
        err.write_text(str(e), encoding='utf-8')
        print('실패', flush=True)


if __name__ == '__main__':
    for d in sys.argv[1:]:
        process(d)
